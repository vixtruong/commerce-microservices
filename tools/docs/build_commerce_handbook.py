"""Build the Vietnamese Commerce communication and observability handbook."""

from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = ROOT / "docs" / "So_tay_giao_tiep_bao_mat_va_quan_sat_Commerce.docx"
ASSET_DIR = ROOT / ".codex-docs" / "commerce-handbook-assets"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DDEBF7"
PALE_BLUE = "EEF5FB"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9E2F3"
TEXT_GRAY = "404040"
WHITE = "FFFFFF"
GREEN = "548235"
ORANGE = "C65911"
RED = "C00000"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    """Load a Windows font that supports Vietnamese glyphs."""
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def set_cell_shading(cell, fill: str) -> None:
    """Apply a background color to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "B7C9DD", size: str = "5") -> None:
    """Draw a subtle border around one table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    """Repeat a table header row when the table spans multiple pages."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_together(paragraph) -> None:
    """Keep a paragraph together and avoid isolated headings."""
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True


def add_page_number(paragraph) -> None:
    """Insert a dynamic PAGE field into a paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """Append a clickable external hyperlink to a paragraph."""
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text(paragraph, text: str, bold: bool = False, italic: bool = False, color: str | None = None):
    """Append consistently styled text to a paragraph."""
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    """Add a bullet paragraph with controlled indentation."""
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(2)
    set_keep_together(paragraph)


def add_numbered(document: Document, text: str, level: int = 0) -> None:
    """Add a numbered instruction paragraph."""
    previous = document.paragraphs[-1] if document.paragraphs else None
    if previous is not None and previous.style.name == "Step Number":
        number = getattr(document, "_commerce_step_number", 0) + 1
    else:
        number = 1
    document._commerce_step_number = number
    paragraph = document.add_paragraph(style="Step Number")
    paragraph.paragraph_format.left_indent = Cm(0.75 + (0.45 * level))
    paragraph.paragraph_format.first_line_indent = Cm(-0.55)
    paragraph.add_run(f"{number}. ").bold = True
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_together(paragraph)


def add_labeled_paragraph(document: Document, label: str, text: str) -> None:
    """Add an emphasized label followed by explanatory prose."""
    paragraph = document.add_paragraph()
    add_text(paragraph, label, bold=True, color=NAVY)
    add_text(paragraph, text)
    set_keep_together(paragraph)


def add_code(document: Document, code: str) -> None:
    """Add a compact monospaced code sample without a surrounding callout box."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F7F7F7")
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run(code.strip("\n"))
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8.3)
    run.font.color.rgb = RGBColor.from_string("202020")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    """Add a professionally formatted data table."""
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if row_index % 2 == 0 else PALE_BLUE
        for column_index, value in enumerate(values):
            cell = cells[column_index]
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(8.7)
            if widths:
                cell.width = Inches(widths[column_index])
        table.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(document: Document, image_path: Path, caption: str, width: float = 6.8) -> None:
    """Insert a centered diagram and caption."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.add_run(caption)
    caption_paragraph.paragraph_format.keep_with_next = True


def wrapped(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, fnt, fill: str, width: int, align: str = "center") -> None:
    """Draw wrapped text inside a diagram region."""
    lines: list[str] = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=width) or [""])
    draw.multiline_text(xy, "\n".join(lines), font=fnt, fill=fill, anchor="mm", align=align, spacing=6)


def node(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, detail: str, fill: str = "EAF2F8", outline: str = NAVY) -> None:
    """Draw a labeled rounded rectangle for a diagram."""
    draw.rounded_rectangle(box, radius=18, fill=f"#{fill}", outline=f"#{outline}", width=4)
    x1, y1, x2, y2 = box
    wrapped(draw, ((x1 + x2) // 2, y1 + 33), title, font(24, True), f"#{NAVY}", 26)
    wrapped(draw, ((x1 + x2) // 2, (y1 + y2) // 2 + 20), detail, font(19), "#303030", 34)


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE, width: int = 5) -> None:
    """Draw a directional arrow between diagram nodes."""
    draw.line([start, end], fill=f"#{color}", width=width)
    x1, y1 = start
    x2, y2 = end
    import math

    angle = math.atan2(y2 - y1, x2 - x1)
    length = 16
    for offset in (2.55, -2.55):
        tip = (x2 - length * math.cos(angle + offset), y2 - length * math.sin(angle + offset))
        draw.line([end, tip], fill=f"#{color}", width=width)


def diagram_architecture(path: Path) -> None:
    """Create the service communication overview diagram."""
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    wrapped(draw, (900, 45), "KIẾN TRÚC GIAO TIẾP COMMERCE", font(30, True), f"#{NAVY}", 60)
    node(draw, (60, 130, 330, 300), "Client", "Postman / Web / Mobile\nREST JSON + JWT", "FCE4D6", ORANGE)
    node(draw, (440, 130, 750, 300), "Commerce.Gateway", "YARP · :8080\nJWT · rate limit · correlation", "DDEBF7")
    services = [
        (840, 100, 1100, 245, "Identity", "JWT + refresh token"),
        (1160, 100, 1420, 245, "Catalog", "Sản phẩm + gRPC"),
        (1480, 100, 1740, 245, "Cart", "Giỏ hàng + Redis"),
        (840, 300, 1100, 445, "Ordering", "Đơn hàng + saga"),
        (1160, 300, 1420, 445, "Inventory", "Tồn kho + reservation"),
        (1480, 300, 1740, 445, "Payment", "Mô phỏng thanh toán"),
        (1000, 500, 1260, 645, "Shipping", "Giao hàng"),
        (1360, 500, 1620, 645, "Notification", "Consumer thông báo"),
    ]
    for x1, y1, x2, y2, title, detail in services:
        node(draw, (x1, y1, x2, y2), title, detail, "E2F0D9", GREEN)
    node(draw, (390, 430, 730, 620), "RabbitMQ", "Topic exchange: commerce.events\nRetry queue · DLX · dead queue", "FFF2CC", ORANGE)
    node(draw, (60, 450, 310, 610), "PostgreSQL", "Mỗi service một database\nKhông truy cập chéo", "F2F2F2", TEXT_GRAY)
    node(draw, (170, 760, 520, 940), "OpenTelemetry Collector", "OTLP 4317/4318\ntrace → Jaeger\nmetric → Prometheus", "E4DFEC", "7030A0")
    node(draw, (650, 760, 930, 940), "Prometheus", "Scrape :8889 và :15692\nPromQL", "FCE4D6", ORANGE)
    node(draw, (1040, 760, 1320, 940), "Grafana", "Dashboard / Explore / Alert", "FFF2CC", ORANGE)
    node(draw, (1430, 760, 1710, 940), "Jaeger", "Trace search / timeline", "DDEBF7", BLUE)
    arrow(draw, (330, 215), (440, 215))
    for y in (170, 365, 570):
        arrow(draw, (750, 215), (835, y))
    arrow(draw, (840, 365), (730, 525), ORANGE)
    arrow(draw, (1100, 365), (730, 525), ORANGE)
    arrow(draw, (1480, 365), (730, 525), ORANGE)
    arrow(draw, (1000, 570), (730, 525), ORANGE)
    arrow(draw, (1360, 570), (730, 525), ORANGE)
    arrow(draw, (1160, 210), (1100, 340), GREEN)
    arrow(draw, (1480, 210), (1100, 340), GREEN)
    arrow(draw, (170, 610), (280, 760), "7030A0")
    arrow(draw, (520, 850), (650, 850), ORANGE)
    arrow(draw, (930, 850), (1040, 850), ORANGE)
    arrow(draw, (520, 880), (1430, 880), BLUE)
    wrapped(draw, (785, 160), "REST", font(18, True), f"#{BLUE}", 10)
    wrapped(draw, (795, 520), "Event", font(18, True), f"#{ORANGE}", 10)
    wrapped(draw, (1235, 270), "gRPC", font(18, True), f"#{GREEN}", 10)
    image.save(path, quality=95)


def diagram_auth(path: Path) -> None:
    """Create the authentication and authorization flow diagram."""
    image = Image.new("RGB", (1800, 880), "white")
    draw = ImageDraw.Draw(image)
    wrapped(draw, (900, 45), "LUỒNG XÁC THỰC VÀ PHÂN QUYỀN", font(30, True), f"#{NAVY}", 60)
    boxes = [
        (80, 150, 380, 330, "1. Đăng nhập", "POST /api/auth/login\nemail + password"),
        (500, 150, 820, 330, "2. Identity", "Kiểm tra ASP.NET Identity\nTạo access + refresh token"),
        (940, 150, 1260, 330, "3. Gateway", "Xác minh chữ ký, iss, aud, exp\nTạo correlation id"),
        (1380, 150, 1720, 330, "4. Service API", "[Authorize] / [Authorize(Roles)]\nĐọc sub và role"),
    ]
    for box in boxes:
        node(draw, box[:4], box[4], box[5], "EAF2F8")
    for x1, x2 in ((380, 500), (820, 940), (1260, 1380)):
        arrow(draw, (x1, 240), (x2, 240))
    node(draw, (180, 490, 620, 730), "Access token · 15 phút", "HS256 trong local\nclaims: sub, email, jti, role\nDùng Bearer cho request API", "E2F0D9", GREEN)
    node(draw, (700, 490, 1120, 730), "Refresh token · 7 ngày", "Chuỗi ngẫu nhiên 64 byte\nDB chỉ lưu SHA-256 hash\nRotation + revoke khi refresh", "FFF2CC", ORANGE)
    node(draw, (1200, 490, 1640, 730), "Quyết định quyền", "Customer: dữ liệu của chính mình\nAdmin: quản trị catalog/stock\nEndpoint công khai: auth + catalog read", "FCE4D6", RED)
    arrow(draw, (650, 330), (400, 490), GREEN)
    arrow(draw, (660, 330), (900, 490), ORANGE)
    arrow(draw, (1500, 330), (1430, 490), RED)
    wrapped(draw, (900, 815), "Không gửi token vào log, metric, trace hoặc dashboard.", font(22, True), f"#{RED}", 70)
    image.save(path, quality=95)


def diagram_saga(path: Path) -> None:
    """Create the checkout saga sequence diagram."""
    image = Image.new("RGB", (1800, 1150), "white")
    draw = ImageDraw.Draw(image)
    wrapped(draw, (900, 42), "CHECKOUT SAGA VÀ EVENTUAL CONSISTENCY", font(30, True), f"#{NAVY}", 60)
    lanes = [(130, "Client"), (410, "Ordering"), (690, "Inventory"), (970, "Payment"), (1250, "Shipping"), (1530, "Notification")]
    for x, title in lanes:
        draw.rounded_rectangle((x - 105, 95, x + 105, 155), radius=14, fill=f"#{NAVY}")
        wrapped(draw, (x, 125), title, font(20, True), "#FFFFFF", 18)
        draw.line((x, 155, x, 1050), fill="#B7C9DD", width=3)
    events = [
        (215, 130, 410, "REST checkout + Idempotency-Key", BLUE),
        (305, 410, 690, "inventory reservation requested", ORANGE),
        (395, 690, 410, "inventory reserved", ORANGE),
        (485, 410, 970, "payment requested", ORANGE),
        (575, 970, 410, "payment succeeded", ORANGE),
        (665, 410, 690, "order paid · confirm stock", ORANGE),
        (755, 410, 1250, "order paid · create shipment", ORANGE),
        (845, 1250, 410, "shipment created", ORANGE),
        (935, 410, 1530, "business events", ORANGE),
        (1020, 130, 410, "GET order để theo dõi trạng thái", BLUE),
    ]
    for y, sx, ex, label, color in events:
        arrow(draw, (sx, y), (ex, y), color)
        wrapped(draw, ((sx + ex) // 2, y - 22), label, font(17, True), f"#{color}", 42)
    wrapped(draw, (900, 1095), "Mỗi bước commit dữ liệu cục bộ + Outbox; consumer dùng Inbox để chống xử lý lặp.", font(21, True), f"#{NAVY}", 80)
    image.save(path, quality=95)


def diagram_rabbitmq(path: Path) -> None:
    """Create the RabbitMQ reliability topology diagram."""
    image = Image.new("RGB", (1800, 950), "white")
    draw = ImageDraw.Draw(image)
    wrapped(draw, (900, 45), "TOPOLOGY RABBITMQ: MAIN · RETRY · DEAD", font(30, True), f"#{NAVY}", 60)
    node(draw, (70, 160, 370, 350), "Outbox publisher", "persistent message\nmandatory publish\npublisher confirm ≤ 10s", "E2F0D9", GREEN)
    node(draw, (500, 160, 850, 350), "commerce.events", "Durable topic exchange\nrouting key *.v1", "DDEBF7", BLUE)
    node(draw, (1010, 120, 1370, 330), "<queue>.v1", "Durable consumer queue\nprefetch 16 · manual ack\nInbox idempotency", "FFF2CC", ORANGE)
    node(draw, (1010, 500, 1370, 700), "<queue>.retry", "TTL 5 giây\nx-retry-count tăng dần\ndead-letter về main exchange", "FCE4D6", ORANGE)
    node(draw, (1480, 500, 1750, 700), "<queue>.dead", "Giữ lỗi cuối cùng\nđể điều tra / replay có kiểm soát", "F4CCCC", RED)
    node(draw, (500, 500, 850, 700), "commerce.events.dlx", "Dead-letter exchange\nDurable topic · binding #", "E4DFEC", "7030A0")
    arrow(draw, (370, 255), (500, 255), GREEN)
    arrow(draw, (850, 255), (1010, 225), BLUE)
    arrow(draw, (1190, 330), (1190, 500), ORANGE)
    arrow(draw, (1010, 600), (850, 600), "7030A0")
    arrow(draw, (680, 500), (680, 350), "7030A0")
    arrow(draw, (1370, 600), (1480, 600), RED)
    wrapped(draw, (1190, 410), "Lỗi tạm thời (< 5 lần)", font(19, True), f"#{ORANGE}", 28)
    wrapped(draw, (1615, 770), "Quá MaxDeliveryAttempts", font(19, True), f"#{RED}", 28)
    wrapped(draw, (900, 860), "ACK chỉ sau khi xử lý/republish thành công; NACK requeue=false khi hết retry.", font(22, True), f"#{NAVY}", 78)
    image.save(path, quality=95)


def diagram_observability(path: Path) -> None:
    """Create the observability data flow diagram."""
    image = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(image)
    wrapped(draw, (900, 45), "PIPELINE QUAN SÁT HỆ THỐNG", font(30, True), f"#{NAVY}", 60)
    node(draw, (60, 170, 410, 390), ".NET services", "ASP.NET Core + HttpClient\nCommerce.Messaging ActivitySource\nRuntime metrics", "E2F0D9", GREEN)
    node(draw, (540, 170, 930, 390), "OTel Collector", "Receiver OTLP gRPC :4317\nProcessor batch\nRouter telemetry", "E4DFEC", "7030A0")
    node(draw, (1080, 90, 1430, 300), "Jaeger", "Nhận trace qua OTLP\nUI :16686\nSearch và critical path", "DDEBF7", BLUE)
    node(draw, (1080, 410, 1430, 620), "Prometheus", "Scrape Collector :8889\nScrape RabbitMQ :15692\nPromQL · UI :9090", "FCE4D6", ORANGE)
    node(draw, (1510, 410, 1760, 620), "Grafana", "Data source Prometheus\nDashboard · Explore · Alert\nUI :3000", "FFF2CC", ORANGE)
    node(draw, (60, 590, 410, 770), "Correlation", "X-Correlation-Id\ntraceparent / tracestate\nBaggage qua HTTP + RabbitMQ", "F2F2F2", TEXT_GRAY)
    arrow(draw, (410, 280), (540, 280), "7030A0")
    arrow(draw, (930, 260), (1080, 195), BLUE)
    arrow(draw, (930, 320), (1080, 515), ORANGE)
    arrow(draw, (1430, 515), (1510, 515), ORANGE)
    arrow(draw, (240, 590), (240, 390), TEXT_GRAY)
    wrapped(draw, (475, 245), "OTLP", font(20, True), "#7030A0", 12)
    wrapped(draw, (1000, 175), "traces", font(18, True), f"#{BLUE}", 12)
    wrapped(draw, (1000, 485), "metrics", font(18, True), f"#{ORANGE}", 12)
    wrapped(draw, (900, 825), "Grafana không nhận telemetry trực tiếp; nó truy vấn Prometheus để hiển thị và cảnh báo.", font(21, True), f"#{NAVY}", 80)
    image.save(path, quality=95)


def configure_document() -> Document:
    """Create the base Word document and configure reusable styles."""
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    section.different_first_page_header_footer = True

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)

    for name, size, before, after in (
        ("Title", 27, 0, 12),
        ("Subtitle", 15, 0, 10),
        ("Heading 1", 18, 14, 7),
        ("Heading 2", 14, 11, 5),
        ("Heading 3", 11.5, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.left_indent = Cm(0)
        style.paragraph_format.first_line_indent = Cm(0)

        # The default Word Title style carries a bottom border; the handbook
        # intentionally uses whitespace instead of a decorative rule.
        if name == "Title":
            paragraph_properties = style.element.get_or_add_pPr()
            paragraph_borders = paragraph_properties.find(qn("w:pBdr"))
            if paragraph_borders is not None:
                paragraph_properties.remove(paragraph_borders)

    styles["Caption"].font.name = "Arial"
    styles["Caption"].font.size = Pt(8.7)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor.from_string(TEXT_GRAY)

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(10.1)

    if "Step Number" not in styles:
        step_style = styles.add_style("Step Number", WD_STYLE_TYPE.PARAGRAPH)
        step_style.font.name = "Arial"
        step_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        step_style.font.size = Pt(10.1)
        step_style.font.color.rgb = RGBColor.from_string(TEXT_GRAY)
        step_style.paragraph_format.space_after = Pt(3)
        step_style.paragraph_format.keep_together = True
        step_style.paragraph_format.widow_control = True

    if "Code Inline" not in styles:
        code_style = styles.add_style("Code Inline", WD_STYLE_TYPE.CHARACTER)
        code_style.font.name = "Consolas"
        code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code_style.font.size = Pt(9)
        code_style.font.color.rgb = RGBColor.from_string("7F2020")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("COMMERCE · SỔ TAY KỸ THUẬT")
    header_run.font.name = "Arial"
    header_run.font.size = Pt(8)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor.from_string(NAVY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Nội bộ · 29/09/2026   |   ")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("666666")
    add_page_number(footer)
    return document


def chapter(document: Document, title: str) -> None:
    """Start a numbered chapter on a fresh page."""
    heading = document.add_heading(title, level=1)
    # A paragraph-level break avoids a blank page when the previous content
    # happens to end exactly at the bottom margin.
    heading.paragraph_format.page_break_before = True


def build_document() -> Path:
    """Generate diagrams and assemble the final handbook."""
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    diagrams = {
        "architecture": ASSET_DIR / "01-architecture.png",
        "auth": ASSET_DIR / "02-auth.png",
        "saga": ASSET_DIR / "03-saga.png",
        "rabbitmq": ASSET_DIR / "04-rabbitmq.png",
        "observability": ASSET_DIR / "05-observability.png",
    }
    diagram_architecture(diagrams["architecture"])
    diagram_auth(diagrams["auth"])
    diagram_saga(diagrams["saga"])
    diagram_rabbitmq(diagrams["rabbitmq"])
    diagram_observability(diagrams["observability"])

    doc = configure_document()

    # Cover
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("SỔ TAY GIAO TIẾP, BẢO MẬT\nVÀ QUAN SÁT HỆ THỐNG COMMERCE")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Kiến trúc microservices .NET · Docker local · Hướng dẫn vận hành và tùy biến")
    doc.add_paragraph()
    cover = doc.add_table(rows=4, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    metadata = [
        ("Phạm vi", "Identity, Gateway, Catalog, Cart, Ordering, Inventory, Payment, Shipping, Notification"),
        ("Hạ tầng", "PostgreSQL, Redis, RabbitMQ, OpenTelemetry Collector, Prometheus, Grafana, Jaeger"),
        ("Cơ sở", "Mã nguồn và cấu hình Docker hiện tại trong repository Commerce"),
        ("Ngày phát hành", "29/09/2026 · Phiên bản 1.0"),
    ]
    for index, (key, value) in enumerate(metadata):
        for cell in cover.rows[index].cells:
            set_cell_border(cell, "D9E2F3")
        set_cell_shading(cover.cell(index, 0), LIGHT_BLUE)
        set_cell_shading(cover.cell(index, 1), WHITE)
        cover.cell(index, 0).width = Inches(1.45)
        cover.cell(index, 1).width = Inches(5.2)
        p1 = cover.cell(index, 0).paragraphs[0]
        p1.add_run(key).bold = True
        p2 = cover.cell(index, 1).paragraphs[0]
        p2.add_run(value)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Tài liệu vận hành dành cho developer, reviewer và người trực hệ thống local.", italic=True, color=TEXT_GRAY)

    # Reader guide and TOC
    doc.add_page_break()
    doc.add_heading("Cách dùng tài liệu", level=1)
    doc.add_paragraph(
        "Tài liệu này mô tả đúng hiện trạng của repository tại thời điểm phát hành, sau đó tách riêng các khuyến nghị để nâng cấp production. "
        "Khi có khác biệt, mã nguồn và các tệp cấu hình được liệt kê ở cuối tài liệu là nguồn sự thật ưu tiên."
    )
    add_labeled_paragraph(doc, "Đọc nhanh: ", "Chương 1–4 để hiểu kiến trúc và checkout; Chương 5–7 để làm việc với bảo mật/RabbitMQ; Chương 8–11 để quan sát; Chương 12 để xử lý sự cố và hardening.")
    add_labeled_paragraph(doc, "Quy ước: ", "“Hiện tại” nghĩa là đã có trong repo. “Khuyến nghị production” không ngụ ý hệ thống hiện tại đã triển khai nội dung đó.")
    add_labeled_paragraph(doc, "An toàn bí mật: ", "Ví dụ chỉ dùng tên biến và giá trị giả. Không chép JWT, mật khẩu, refresh token hoặc connection string thật vào tài liệu, dashboard hay log.")
    doc.add_heading("Mục lục", level=2)
    toc_items = [
        "1. Bức tranh tổng thể và nguyên tắc biên service",
        "2. Bản đồ service, dữ liệu và cổng local",
        "3. REST, gRPC và event: chọn đúng kiểu giao tiếp",
        "4. Luồng checkout saga từ đầu đến cuối",
        "5. Xác thực bằng ASP.NET Core Identity và JWT",
        "6. Phân quyền, ownership và kiểm soát tại Gateway/service",
        "7. RabbitMQ: cấu hình, độ tin cậy, retry, DLQ và tùy biến",
        "8. OpenTelemetry Collector: thu nhận và định tuyến telemetry",
        "9. Prometheus: scrape, PromQL, recording rule và alert rule",
        "10. Grafana: dashboard, Explore, biến, cảnh báo và provisioning",
        "11. Jaeger: tìm trace, đọc critical path và tùy biến",
        "12. Runbook xử lý sự cố và hardening production",
        "Phụ lục A–D. Lệnh nhanh, ma trận routing, cấu hình và nguồn tham khảo",
    ]
    for item in toc_items:
        add_bullet(doc, item)

    chapter(doc, "1. Bức tranh tổng thể và nguyên tắc biên service")
    doc.add_paragraph(
        "Commerce là hệ thống microservices .NET chạy local bằng Docker Compose. Người dùng chỉ gọi REST qua Commerce.Gateway. "
        "Các service không đọc database của nhau; truy vấn đồng bộ nội bộ dùng gRPC, còn thay đổi nghiệp vụ dài và cần chống lỗi dùng RabbitMQ."
    )
    add_figure(doc, diagrams["architecture"], "Hình 1 — Toàn cảnh các kênh giao tiếp và pipeline quan sát.")
    doc.add_heading("1.1 Bốn nguyên tắc cốt lõi", level=2)
    principles = [
        ("Một cửa công khai", "Client đi qua Gateway/YARP tại http://localhost:8080. Service API và gRPC dùng mạng Docker nội bộ."),
        ("Sở hữu dữ liệu", "Mỗi service có database PostgreSQL riêng; Cart sở hữu Redis. Không join hoặc ghi chéo database."),
        ("Đồng bộ có giới hạn", "gRPC chỉ phục vụ truy vấn dữ liệu cần ngay, có deadline 3 giây; không biến gRPC thành giao dịch phân tán."),
        ("Bất đồng bộ có bằng chứng", "Event được ghi Outbox cùng transaction nghiệp vụ, publish có confirm, consumer có Inbox chống xử lý lặp."),
    ]
    add_table(doc, ["Nguyên tắc", "Ý nghĩa thực tế"], [[a, b] for a, b in principles], [1.55, 5.25])
    doc.add_heading("1.2 Điều gì không được làm", level=2)
    for text in [
        "Không cho Gateway chứa domain logic hoặc truy cập database của service.",
        "Không công khai gRPC, RabbitMQ hay PostgreSQL ra Internet chỉ để tiện debug.",
        "Không dùng một event mutable cho nhiều phiên bản consumer; thay đổi breaking phải phát hành routing key/contract v2.",
        "Không coi HTTP 202 của checkout là đơn hàng đã hoàn tất; đó chỉ là xác nhận saga đã được chấp nhận.",
        "Không lưu token, mật khẩu, PII hoặc connection string trong span attribute, metric label hay log message.",
    ]:
        add_bullet(doc, text)

    chapter(doc, "2. Bản đồ service, dữ liệu và cổng local")
    doc.add_heading("2.1 Trách nhiệm và kho dữ liệu", level=2)
    service_rows = [
        ["Commerce.Gateway", "YARP routing, JWT validation, rate limit, correlation", "Không có database"],
        ["Identity.Api", "Đăng ký, đăng nhập, refresh, logout, phát hành JWT", "commerce_identity"],
        ["Catalog.Api ×2", "Sản phẩm; REST công khai đọc; gRPC nội bộ", "commerce_catalog"],
        ["Cart.Api", "Giỏ hàng theo customer", "Redis"],
        ["Ordering.Api", "Checkout, Order aggregate, saga, outbox/inbox", "commerce_ordering"],
        ["Inventory.Api", "Tồn kho, reservation lease, nhận hàng", "commerce_inventory"],
        ["Payment.Api", "Mô phỏng kết quả thanh toán", "commerce_payment"],
        ["Shipping.Api", "Tạo và cập nhật shipment", "commerce_shipping"],
        ["Notification.Worker", "Tiêu thụ event và lưu thông báo", "commerce_notification"],
    ]
    add_table(doc, ["Thành phần", "Trách nhiệm", "Kho sở hữu"], service_rows, [1.45, 3.65, 1.65])
    doc.add_heading("2.2 Cổng từ máy host", level=2)
    port_rows = [
        ["Gateway", "8080", "Public REST entry point"],
        ["PostgreSQL", "127.0.0.1:${POSTGRES_PORT:-5432}", "DB management local; bind loopback"],
        ["Redis", "127.0.0.1:${REDIS_PORT:-6379}", "Redis management local; hiện không có auth"],
        ["RabbitMQ Management", "15672", "UI; AMQP 5672 chỉ ở mạng Docker"],
        ["Jaeger", "16686", "Trace UI"],
        ["Prometheus", "9090", "Query/Targets UI"],
        ["Grafana", "3000", "Dashboard/Explore/Alert"],
        ["OTel Collector", "không publish ra host", "OTLP 4317/4318 và exporter 8889 chỉ nội bộ"],
    ]
    add_table(doc, ["Thành phần", "Cổng local", "Mục đích"], port_rows, [1.55, 2.35, 2.85])
    add_labeled_paragraph(doc, "Lưu ý mạng: ", "Tên host trong container là tên service Docker như postgres, redis, rabbitmq, otel-collector. Từ Navicat/Redis client trên Windows dùng 127.0.0.1 và cổng đã publish.")
    doc.add_heading("2.3 Khởi động và kiểm tra nhanh", level=2)
    add_code(doc, """
Copy-Item .env.example .env
docker compose --env-file .env up -d --build
docker compose --env-file .env ps
docker compose --env-file .env logs --tail 100 gateway identity-api ordering-api
""")
    doc.add_paragraph(
        "Giữ `.env` ngoài Git. `.env.example` chỉ là mẫu local. Nếu thay cổng Redis bằng `REDIS_PORT`, thêm biến đó vào `.env`; nếu bỏ trống, Compose dùng 6379."
    )

    chapter(doc, "3. REST, gRPC và event: chọn đúng kiểu giao tiếp")
    doc.add_heading("3.1 REST qua Gateway", level=2)
    doc.add_paragraph(
        "REST phù hợp với request/response do người dùng khởi tạo. Gateway route theo prefix `/api/...`, chuyển tiếp header Authorization và correlation, "
        "đồng thời áp dụng token-bucket rate limit theo claim `sub` hoặc địa chỉ IP. Cấu hình hiện tại: 100 token, nạp 50 token mỗi 10 giây, queue tối đa 10; vượt ngưỡng trả 429."
    )
    rest_rows = [
        ["/api/auth/*", "Identity.Api", "register, login, refresh, logout"],
        ["/api/catalog/*", "Catalog.Api (round-robin 2 instance)", "đọc sản phẩm; Admin quản trị"],
        ["/api/cart/*", "Cart.Api", "giỏ hàng của customer"],
        ["/api/orders/*", "Ordering.Api", "checkout và đọc đơn"],
        ["/api/inventory/*", "Inventory.Api", "Admin xem/nhập tồn"],
        ["/api/payments/*", "Payment.Api", "Admin tra thanh toán"],
        ["/api/shipping/*", "Shipping.Api", "tra shipment; Admin advance"],
    ]
    add_table(doc, ["Route", "Đích", "Ngữ nghĩa"], rest_rows, [1.6, 2.35, 2.8])
    doc.add_heading("3.2 gRPC nội bộ", level=2)
    grpc_rows = [
        ["Cart → Catalog", "GetProduct", "Kiểm tra snapshot sản phẩm khi cập nhật giỏ", "3 giây"],
        ["Ordering → Cart", "GetCartForCheckout", "Lấy giỏ hiện tại khi checkout", "3 giây"],
        ["Ordering → Catalog", "GetProducts", "Lấy nhiều snapshot sản phẩm", "3 giây"],
    ]
    add_table(doc, ["Caller", "RPC", "Mục đích", "Deadline"], grpc_rows, [1.45, 1.35, 3.2, 0.75])
    doc.add_paragraph(
        "gRPC chạy HTTP/2 ở cổng container 8081. Deadline ngăn thread treo vô hạn, nhưng caller vẫn phải xử lý timeout như một lỗi có thể retry ở biên phù hợp. "
        "Không dùng gRPC để cập nhật đồng thời hai database."
    )
    doc.add_heading("3.3 RabbitMQ cho business event", level=2)
    doc.add_paragraph(
        "RabbitMQ phù hợp khi publisher không cần đợi toàn bộ chuỗi nghiệp vụ hoàn thành. Routing key chứa capability, sự kiện và version, ví dụ "
        "`ordering.payment-requested.v1`. Consumer queue có tên riêng nên mỗi service nhận bản sao độc lập."
    )
    add_table(
        doc,
        ["Câu hỏi", "REST", "gRPC", "RabbitMQ"],
        [
            ["Cần kết quả ngay?", "Có", "Có", "Không"],
            ["Caller là client?", "Có", "Không", "Không"],
            ["Ghép lỏng / fan-out?", "Thấp", "Thấp", "Cao"],
            ["Lỗi downstream?", "Trả lỗi/timeout", "Status/timeout", "Retry + DLQ"],
            ["Nhất quán", "Trong request", "Trong request", "Eventual consistency"],
        ],
        [1.55, 1.5, 1.5, 2.2],
    )
    doc.add_heading("3.4 Correlation và trace propagation", level=2)
    doc.add_paragraph(
        "Middleware đọc `X-Correlation-Id`; nếu thiếu hoặc không hợp lệ, hệ thống tạo GUID mới và trả header này trong response. Với HTTP, OpenTelemetry truyền W3C `traceparent`/`tracestate`. "
        "Với RabbitMQ, publisher ghi các header trace vào message; consumer tạo span mới nhưng giữ quan hệ trace để Jaeger hiển thị chuỗi end-to-end."
    )
    add_code(doc, """
$headers = @{
  Authorization      = "Bearer <access-token>"
  "X-Correlation-Id" = [guid]::NewGuid().ToString()
  "Idempotency-Key"  = [guid]::NewGuid().ToString()
}
Invoke-RestMethod -Method Post -Uri http://localhost:8080/api/orders/checkout `
  -Headers $headers -ContentType "application/json" -Body "{}"
""")

    chapter(doc, "4. Luồng checkout saga từ đầu đến cuối")
    add_figure(doc, diagrams["saga"], "Hình 2 — Happy path của checkout saga; các mũi tên màu cam là event RabbitMQ.")
    doc.add_heading("4.1 Happy path", level=2)
    steps = [
        "Client gửi checkout có Bearer token và `Idempotency-Key`.",
        "Ordering xác định customer từ claim `sub`, lấy cart qua gRPC, lấy snapshot sản phẩm qua Catalog gRPC.",
        "Ordering commit Order, SagaState và Outbox message trong một transaction; API trả 202 Accepted.",
        "Outbox publisher gửi `ordering.inventory-reservation-requested.v1` và chỉ đánh dấu processed sau publisher confirm.",
        "Inventory tạo reservation rồi phát `inventory.reserved.v1`; nếu thiếu hàng phát `inventory.reservation-failed.v1`.",
        "Ordering nhận kết quả tồn kho. Nếu thành công, phát `ordering.payment-requested.v1`.",
        "Payment phát `payment.succeeded.v1` hoặc `payment.failed.v1`.",
        "Nếu thanh toán thành công, Ordering chuyển Order sang Paid và phát `ordering.order-paid.v1`; Inventory xác nhận stock, Shipping tạo shipment.",
        "Notification tiêu thụ các event quan trọng. Client gọi GET order để thấy trạng thái cuối cùng.",
    ]
    for step in steps:
        add_numbered(doc, step)
    doc.add_heading("4.2 Nhánh lỗi và bù trừ", level=2)
    add_bullet(doc, "Reservation thất bại hoặc hết hạn: saga hủy đơn; không gửi payment request.")
    add_bullet(doc, "Payment thất bại: Ordering phát inventory release request và order cancelled.")
    add_bullet(doc, "Consumer lỗi tạm thời: message đi retry queue 5 giây rồi quay về queue chính.")
    add_bullet(doc, "Lỗi quá số lần: message vào dead queue; operator phải điều tra nguyên nhân trước khi replay.")
    add_bullet(doc, "Event phát trùng: Inbox ngăn mutation lặp theo cặp MessageId + Consumer.")
    doc.add_heading("4.3 Idempotency khác retry như thế nào", level=2)
    doc.add_paragraph(
        "Retry đảm bảo thao tác có cơ hội chạy lại; idempotency đảm bảo chạy lại không tạo kết quả nghiệp vụ lặp. `Idempotency-Key` bảo vệ lệnh checkout từ client. "
        "MessageId/Inbox bảo vệ consumer. Outbox bảo vệ khoảng hở giữa commit database và publish broker. Ba cơ chế giải quyết ba điểm lỗi khác nhau và phải tồn tại đồng thời."
    )
    doc.add_heading("4.4 Cách theo dõi một checkout", level=2)
    add_numbered(doc, "Ghi lại `X-Correlation-Id`, order id và thời điểm gửi request.")
    add_numbered(doc, "Trong Jaeger, chọn service `Commerce.Gateway` hoặc `Ordering.Api`, lookback phù hợp, lọc operation/request path nếu cần.")
    add_numbered(doc, "Trong RabbitMQ UI, kiểm tra queue Ready/Unacked/Consumers và dead queue của consumer bị nghi ngờ.")
    add_numbered(doc, "Trong Prometheus/Grafana, đối chiếu request rate, latency và RabbitMQ backlog đúng khoảng thời gian.")
    add_numbered(doc, "Chỉ kiểm tra DB của service sở hữu dữ liệu; không sửa trạng thái trực tiếp để “chữa” saga.")

    chapter(doc, "5. Xác thực bằng ASP.NET Core Identity và JWT")
    add_figure(doc, diagrams["auth"], "Hình 3 — Access token dùng cho API; refresh token chỉ dùng để xin access token mới.")
    doc.add_heading("5.1 Cách Identity tạo token", level=2)
    doc.add_paragraph(
        "Identity.Api dùng ASP.NET Core Identity để quản lý user, password hash, lockout và role. Password policy local yêu cầu tối thiểu 12 ký tự, chữ hoa, chữ thường, chữ số và ký tự đặc biệt; email là duy nhất; lockout sau tối đa 5 lần thất bại."
    )
    jwt_rows = [
        ["Access token", "15 phút", "JWT HS256", "sub, email, jti, role", "Authorization: Bearer"],
        ["Refresh token", "7 ngày", "64 byte ngẫu nhiên", "DB chỉ lưu SHA-256 hash", "POST /api/auth/refresh"],
    ]
    add_table(doc, ["Token", "Thời hạn", "Hình thức", "Nội dung/lưu trữ", "Cách dùng"], jwt_rows, [1.0, 0.8, 1.15, 2.15, 1.65])
    doc.add_paragraph(
        "Khi refresh, token cũ bị revoke và liên kết tới hash của token thay thế. Logout revoke refresh token. Access token đã phát không được thu hồi tức thì; thời hạn ngắn giới hạn cửa sổ rủi ro."
    )
    doc.add_heading("5.2 Secret key nằm ở đâu và chạy qua các lớp config", level=2)
    add_table(
        doc,
        ["Lớp", "Tên", "Vai trò"],
        [
            ["`.env` local", "JWT_SIGNING_KEY", "Giá trị bí mật thật trên máy developer; không commit"],
            ["docker-compose.yml", "Jwt__SigningKey", "Map environment variable vào ASP.NET Core configuration"],
            [".NET configuration", "Jwt:SigningKey", "Identity dùng để ký; Gateway và API dùng để xác minh"],
            ["Mã nguồn", "AddCommerceJwt / token service", "Validate issuer, audience, lifetime, signature; Identity ký HS256"],
        ],
        [1.3, 1.75, 3.7],
    )
    add_code(doc, """
# .env — ví dụ, không dùng giá trị này ngoài local
JWT_ISSUER=commerce.identity
JWT_AUDIENCE=commerce.apis
JWT_SIGNING_KEY=<chuoi-ngau-nhien-dai-va-bi-mat>

# Docker Compose chuyển thành hierarchical configuration
Jwt__Issuer=${JWT_ISSUER}
Jwt__Audience=${JWT_AUDIENCE}
Jwt__SigningKey=${JWT_SIGNING_KEY}
""")
    add_labeled_paragraph(doc, "Vì sao mọi API cần key? ", "Với HS256, cùng một secret vừa ký vừa xác minh. Vì vậy hiện tại Identity, Gateway và các API đều nhận key. Đây là thuận tiện local nhưng mở rộng blast radius nếu một service bị lộ.")
    add_labeled_paragraph(doc, "Khuyến nghị production: ", "Dùng OIDC/OAuth 2.0 với RS256/ES256 và JWKS. Chỉ Identity/authorization server giữ private key; Gateway và service chỉ nhận public key. Đưa secret/private key vào secret manager và xoay khóa có `kid`.")
    doc.add_heading("5.3 Thử đăng nhập và gọi API", level=2)
    add_code(doc, """
$loginBody = @{
  email    = "<customer-email-from-.env>"
  password = "<customer-password-from-.env>"
} | ConvertTo-Json

$tokens = Invoke-RestMethod -Method Post `
  -Uri http://localhost:8080/api/auth/login `
  -ContentType "application/json" -Body $loginBody

Invoke-RestMethod -Uri http://localhost:8080/api/cart `
  -Headers @{ Authorization = "Bearer $($tokens.accessToken)" }
""")
    doc.add_paragraph(
        "Không paste access token vào issue, screenshot hoặc query URL. Nếu cần debug claim, dùng công cụ nội bộ/offline và che chữ ký/token trước khi chia sẻ."
    )
    doc.add_heading("5.4 Gateway và service cùng validate", level=2)
    doc.add_paragraph(
        "Gateway validate sớm để chặn token sai và chuẩn hóa security boundary. Service tiếp tục validate `[Authorize]` vì mạng nội bộ không phải là trust boundary tuyệt đối. "
        "Cấu hình hiện tại tắt inbound claim mapping, dùng claim `role` làm role claim, và cho phép clock skew 30 giây."
    )

    chapter(doc, "6. Phân quyền, ownership và kiểm soát tại Gateway/service")
    doc.add_heading("6.1 Ma trận quyền hiện tại", level=2)
    auth_rows = [
        ["Auth register/login/refresh", "Anonymous", "Identity xử lý credential/token"],
        ["Catalog GET", "Anonymous", "Public read"],
        ["Catalog create/update/activate/deactivate", "Admin", "Role claim"],
        ["Cart GET/PUT/DELETE", "Authenticated", "Customer id lấy từ `sub`"],
        ["Order checkout/read", "Authenticated", "Customer-scoped bằng `sub`"],
        ["Inventory GET/receipts", "Admin", "Role claim"],
        ["Payment GET by order", "Admin", "Role claim"],
        ["Shipping GET", "Authenticated", "Hiện chưa kiểm tra owner trong action"],
        ["Shipping advance", "Admin", "Role claim"],
    ]
    add_table(doc, ["Nhóm endpoint", "Quyền", "Cách enforce/ghi chú"], auth_rows, [2.4, 1.25, 3.1])
    add_labeled_paragraph(doc, "Điểm cần review: ", "Shipping GET hiện yêu cầu authenticated nhưng chưa chứng minh shipment thuộc customer đang gọi. Trước production nên thêm ownership check hoặc chỉ cho Admin, tùy use case.")
    doc.add_heading("6.2 Authentication khác authorization", level=2)
    add_bullet(doc, "Authentication trả lời “ai đang gọi?” bằng chữ ký JWT và claim `sub`.")
    add_bullet(doc, "Role authorization trả lời “người này có vai trò nào?” bằng claim `role`.")
    add_bullet(doc, "Resource authorization trả lời “người này có sở hữu order/shipment này không?” bằng so sánh `sub` với dữ liệu thuộc service.")
    add_bullet(doc, "Business authorization trả lời “trạng thái hiện tại có cho phép thao tác không?” trong domain/application layer.")
    doc.add_heading("6.3 Cách thêm policy tùy biến", level=2)
    doc.add_paragraph(
        "Khi rule vượt quá role đơn giản, tạo policy/authorization handler ở service sở hữu resource. Ví dụ policy `CanReadOrder` đọc `sub`, tải order bằng repository của Ordering, rồi xác nhận CustomerId trùng. "
        "Không nhét rule này vào Gateway vì Gateway không sở hữu order."
    )
    add_code(doc, """
builder.Services.AddAuthorization(options =>
{
    options.AddPolicy("CatalogWrite", policy =>
        policy.RequireRole("Admin")
              .RequireClaim("scope", "catalog.write"));
});

// Endpoint vẫn phải kiểm tra ownership nếu resource gắn với customer.
""")
    doc.add_heading("6.4 Checklist thêm endpoint", level=2)
    for item in [
        "Quyết định endpoint public, authenticated, role-based hay resource-based.",
        "Khai báo `[Authorize]`, role/policy tại service; không chỉ dựa vào route Gateway.",
        "Không nhận CustomerId từ body nếu đã có `sub`; tránh IDOR/mạo danh.",
        "Trả 401 khi thiếu/sai token, 403 khi đã xác thực nhưng không đủ quyền, 404 khi cần tránh lộ sự tồn tại resource.",
        "Cập nhật Postman collection khi REST contract/auth workflow thay đổi.",
        "Thêm test cho anonymous, sai role, đúng role và sai ownership.",
    ]:
        add_bullet(doc, item)

    chapter(doc, "7. RabbitMQ: cấu hình, độ tin cậy, retry, DLQ và tùy biến")
    add_figure(doc, diagrams["rabbitmq"], "Hình 4 — Đường đi của message khi thành công, retry và lỗi cuối cùng.")
    doc.add_heading("7.1 Cấu hình hiện tại", level=2)
    add_table(
        doc,
        ["Khóa", "Giá trị/nguồn", "Ý nghĩa"],
        [
            ["RabbitMq__ConnectionString", "amqp://<user>:<pass>@rabbitmq:5672", "AMQP nội bộ Docker"],
            ["RabbitMq__ExchangeName", "commerce.events", "Durable topic exchange chính"],
            ["RabbitMq__DeadLetterExchangeName", "commerce.events.dlx", "Durable topic DLX"],
            ["MaxDeliveryAttempts", "5", "Số lần thử trước khi dead-letter cuối"],
            ["Retry delay", "5.000 ms", "TTL của queue retry"],
            ["Prefetch", "16", "Tối đa message chưa ACK trên channel/consumer"],
            ["Publisher confirm timeout", "10 giây", "Publisher chỉ hoàn tất khi broker confirm"],
        ],
        [1.9, 2.2, 2.55],
    )
    doc.add_heading("7.2 Envelope và trace context", level=2)
    doc.add_paragraph(
        "Publisher dùng message persistent, `mandatory=true` và publisher confirms. Message mang MessageId, CorrelationId, Type/routing key, JSON content, cùng `traceparent`, `tracestate` và causation id. "
        "Schema contract không phải domain entity; chỉ chứa dữ liệu cần để service khác phản ứng."
    )
    doc.add_heading("7.3 Publisher: Outbox", level=2)
    doc.add_paragraph(
        "Outbox message được ghi trong cùng local transaction với Order/Reservation/Payment. Worker đọc theo batch 20, claim lease 30 giây, poll khoảng 2 giây khi rỗng và nhanh hơn khi còn việc. "
        "Publisher thử tối đa 10 lần; chỉ đánh dấu processed sau khi broker confirm. Vì crash có thể xảy ra sau publish nhưng trước mark-processed, consumer vẫn phải idempotent."
    )
    doc.add_heading("7.4 Consumer: ACK, retry và Inbox", level=2)
    add_numbered(doc, "Broker giao message, consumer dùng manual ACK và prefetch 16.")
    add_numbered(doc, "Consumer mở local transaction, kiểm tra Inbox theo MessageId + Consumer.")
    add_numbered(doc, "Nếu chưa xử lý, thực hiện mutation và ghi Inbox trong cùng transaction.")
    add_numbered(doc, "Thành công thì ACK. Lỗi tạm thời thì publish sang retry queue, tăng `x-retry-count`, sau đó ACK bản gốc.")
    add_numbered(doc, "Retry queue hết TTL 5 giây sẽ dead-letter về exchange chính/routing key ban đầu.")
    add_numbered(doc, "Đủ 5 lần vẫn lỗi thì NACK `requeue=false`; broker đưa vào DLX và `<queue>.dead`.")
    doc.add_heading("7.5 Routing key và queue quan trọng", level=2)
    routing_rows = [
        ["ordering.inventory-reservation-requested.v1", "inventory.reserve-order.v1", "Inventory"],
        ["inventory.reserved.v1", "ordering.inventory-reserved.v1", "Ordering"],
        ["inventory.reservation-failed.v1", "ordering.inventory-failed.v1", "Ordering"],
        ["ordering.payment-requested.v1", "payment.process-order.v1", "Payment"],
        ["payment.succeeded.v1", "ordering.payment-succeeded.v1", "Ordering"],
        ["payment.failed.v1", "ordering.payment-failed.v1", "Ordering"],
        ["ordering.order-paid.v1", "inventory.confirm-paid-order.v1", "Inventory"],
        ["ordering.order-paid.v1", "shipping.create-paid-order.v1", "Shipping"],
        ["shipping.shipment-created.v1", "ordering.shipment-created.v1", "Ordering"],
        ["*.order/payment/shipment events", "notification.*.v1", "Notification"],
    ]
    add_table(doc, ["Routing key", "Queue", "Consumer"], routing_rows, [3.2, 2.45, 1.1])
    doc.add_heading("7.6 Dùng RabbitMQ Management UI", level=2)
    for instruction in [
        "Mở http://localhost:15672 và đăng nhập bằng `RABBITMQ_DEFAULT_USER/PASS` trong `.env`.",
        "Overview: xem message rate, connection/channel count và node alarm. Alarm memory/disk có thể làm publisher bị flow control.",
        "Exchanges: mở `commerce.events`, xem type=topic, durability và bindings; mở DLX để kiểm tra dead bindings.",
        "Queues and Streams: tìm queue consumer; theo dõi Ready, Unacked, Total, Consumers, Incoming và Deliver/Ack.",
        "Connections/Channels: xác nhận từng container còn kết nối; Unacked tăng mà ack rate đứng yên thường là consumer treo/chậm.",
        "Chỉ dùng Get messages khi debug local. Chọn requeue=true nếu không muốn lấy message khỏi queue.",
    ]:
        add_numbered(doc, instruction)
    doc.add_heading("7.7 Replay dead message an toàn", level=2)
    for item in [
        "Chụp MessageId, routing key, correlation id, headers và exception/log liên quan.",
        "Sửa nguyên nhân gốc trước: schema, dependency, dữ liệu hoặc bug handler.",
        "Kiểm tra handler có Inbox/idempotency và side effect ngoài DB (email/payment) có idempotency riêng.",
        "Republish vào `commerce.events` bằng routing key gốc, giữ MessageId/correlation nếu tool cho phép; không publish vào queue retry tùy tiện.",
        "Theo dõi ACK, Inbox và trạng thái nghiệp vụ. Chỉ sau đó mới purge/delete dead message nếu cần.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("7.8 Thêm event hoặc consumer mới", level=2)
    for item in [
        "Định nghĩa integration contract độc lập domain entity; đặt tên sự kiện ở thì quá khứ nếu là fact, hoặc `*-requested` nếu là command-like event.",
        "Chọn routing key `<capability>.<event>.v1`; breaking change tạo `.v2`, không sửa nghĩa `.v1`.",
        "Đăng ký durable queue riêng của consumer, main binding, retry queue, dead queue và Inbox identity.",
        "Publisher ghi Outbox trong cùng transaction. Consumer ACK sau commit; propagate cancellation và trace context.",
        "Test duplicate delivery, poison message, broker restart, consumer restart và thứ tự event. Không giả định global ordering giữa nhiều queue.",
        "Cập nhật tài liệu/ma trận contract và dashboard backlog/dead queue.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("7.9 Hardening production", level=2)
    add_bullet(doc, "Tạo vhost và user riêng cho từng service; cấp configure/write/read đúng exchange/queue cần thiết, không dùng shared admin account.")
    add_bullet(doc, "Bật TLS, quản lý certificate/credential bằng secret manager, xoay credential và không publish AMQP port công khai.")
    add_bullet(doc, "Đặt queue length/TTL/policy theo tải, quorum queues nếu cần HA, cảnh báo dead queue và consumer lag.")
    add_bullet(doc, "Chỉ replay tự động khi có invariant rõ; với payment/email nên ưu tiên quy trình phê duyệt.")
    doc.add_heading("7.10 Kiểm thử failure có chủ đích", level=2)
    for item in [
        "Dừng consumer sau khi publisher đã ghi Outbox; khởi động lại và xác nhận message vẫn được xử lý.",
        "Gây lỗi handler tạm thời; xác nhận message đi qua retry queue, `x-retry-count` tăng và không tạo mutation trùng.",
        "Gây poison message; xác nhận message vào `.dead` sau giới hạn retry và queue chính tiếp tục phục vụ message khác.",
        "Restart RabbitMQ; xác nhận durable exchange/queue và persistent message còn tồn tại, publisher tự kết nối lại.",
        "Gửi lại cùng MessageId; xác nhận Inbox bỏ qua side effect lần hai.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("7.11 Chỉ số nên cảnh báo", level=2)
    add_bullet(doc, "Ready tăng liên tục: consumer không theo kịp hoặc đã mất kết nối.")
    add_bullet(doc, "Unacked cao kéo dài: handler chậm/treo, prefetch quá lớn hoặc downstream bị nghẽn.")
    add_bullet(doc, "Consumers = 0 trên queue nghiệp vụ: container/registration gặp lỗi.")
    add_bullet(doc, "Dead queue > 0: cần ticket/điều tra ngay; không chỉ purge để làm biểu đồ xanh.")

    chapter(doc, "8. OpenTelemetry Collector: thu nhận và định tuyến telemetry")
    add_figure(doc, diagrams["observability"], "Hình 5 — Trace và metric dùng chung OTLP receiver nhưng đi tới backend khác nhau.")
    doc.add_heading("8.1 Collector dùng để làm gì", level=2)
    doc.add_paragraph(
        "OpenTelemetry Collector là agent/gateway trung lập vendor. Service gửi telemetry một lần bằng OTLP; Collector nhận, xử lý theo batch và xuất tới backend. "
        "Collector không có dashboard người dùng cuối. Trong stack hiện tại, trace đi tới Jaeger; metric được chuyển thành endpoint Prometheus scrape."
    )
    doc.add_heading("8.2 Pipeline hiện tại", level=2)
    add_code(doc, """
receivers:
  otlp:
    protocols:
      grpc: { endpoint: 0.0.0.0:4317 }
      http: { endpoint: 0.0.0.0:4318 }
processors:
  batch: {}
exporters:
  otlp_grpc/jaeger:
    endpoint: jaeger:4317
    tls: { insecure: true }
  prometheus:
    endpoint: 0.0.0.0:8889
service:
  pipelines:
    traces:  { receivers: [otlp], processors: [batch], exporters: [otlp_grpc/jaeger] }
    metrics: { receivers: [otlp], processors: [batch], exporters: [prometheus] }
""")
    doc.add_paragraph(
        "Mỗi .NET service đặt `service.name` bằng application name, instrument ASP.NET Core/HttpClient/runtime và ActivitySource `Commerce.Messaging`. Docker cung cấp "
        "`OTEL_EXPORTER_OTLP_ENDPOINT=http://otel-collector:4317`, protocol gRPC và resource attribute `deployment.environment=local`."
    )
    doc.add_heading("8.3 Kiểm tra Collector", level=2)
    add_code(doc, """
docker compose --env-file .env ps otel-collector jaeger prometheus
docker compose --env-file .env logs --tail 200 otel-collector

# Từ host, Prometheus phải thấy target commerce-otel là UP:
Invoke-RestMethod http://localhost:9090/api/v1/targets
""")
    add_bullet(doc, "`connection refused` tới jaeger:4317: kiểm tra network/service name và Jaeger readiness.")
    add_bullet(doc, "Metric target down: kiểm tra Collector còn listen 8889 và Prometheus cùng network.")
    add_bullet(doc, "Có HTTP request nhưng không có trace: kiểm tra OTLP endpoint/protocol, sampling và log Collector.")
    doc.add_heading("8.4 Tùy biến an toàn", level=2)
    doc.add_paragraph(
        "Khi tải tăng, thêm `memory_limiter` trước `batch`; khi cần chuẩn hóa metadata, thêm `resource`/`attributes` processor; khi điều tra local, có thể thêm debug exporter tạm thời. "
        "Mỗi processor chỉ có hiệu lực khi được đưa vào đúng pipeline."
    )
    add_code(doc, """
processors:
  memory_limiter:
    check_interval: 1s
    limit_mib: 256
  batch:
    timeout: 5s
    send_batch_size: 1024
  attributes/redact:
    actions:
      - key: http.request.header.authorization
        action: delete

service:
  pipelines:
    traces:
      receivers: [otlp]
      processors: [memory_limiter, attributes/redact, batch]
      exporters: [otlp_grpc/jaeger]
""")
    add_labeled_paragraph(doc, "Sau khi sửa: ", "Chạy `docker compose config`, restart Collector, đọc log lỗi cấu hình, tạo request thử, kiểm tra Jaeger/Prometheus. Không gửi PII/high-cardinality ID vào metric label.")
    doc.add_heading("8.5 Sampling", level=2)
    doc.add_paragraph(
        "Local nên giữ 100% trace để học/debug. Production có thể dùng parent-based ratio sampling tại SDK để giảm dữ liệu, hoặc tail sampling tại Collector nếu cần giữ toàn bộ trace lỗi/chậm. "
        "Tail sampling cần Collector đủ memory và kiến trúc phân phối phù hợp; không bật chỉ bằng một dòng mà không capacity test."
    )

    chapter(doc, "9. Prometheus: scrape, PromQL, recording rule và alert rule")
    doc.add_heading("9.1 Prometheus dùng để làm gì", level=2)
    doc.add_paragraph(
        "Prometheus kéo metric theo chu kỳ, lưu chuỗi thời gian và truy vấn bằng PromQL. Cấu hình hiện tại scrape mỗi 10 giây từ Collector `otel-collector:8889` và RabbitMQ `rabbitmq:15692`. "
        "Do exporter cần tránh xung đột label `job`, `service.name` xuất hiện dưới label `exported_job` trong các truy vấn dashboard hiện tại."
    )
    doc.add_heading("9.2 Quy trình sử dụng UI", level=2)
    for item in [
        "Mở http://localhost:9090/targets. Hai scrape pool `commerce-otel` và `rabbitmq` phải UP.",
        "Mở Graph/Query; bắt đầu với `up`, chuyển Table để xem labels trước khi dựng biểu đồ.",
        "Thu hẹp theo `exported_job`, `http_route` hoặc label RabbitMQ; tránh query không giới hạn trên khoảng thời gian dài.",
        "Dùng Range query và chọn time window; so sánh rate thay vì raw counter.",
        "Nếu query đúng, chuyển sang Grafana để tạo panel lâu dài; Prometheus UI phù hợp điều tra nhanh.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("9.3 PromQL thực dụng", level=2)
    queries = [
        ("Target health", "up"),
        ("Request/giây theo service", "sum by (exported_job) (rate(http_server_request_duration_seconds_count[5m]))"),
        ("P95 latency theo service", "histogram_quantile(0.95, sum by (le, exported_job) (rate(http_server_request_duration_seconds_bucket[5m])))"),
        ("P95 theo route", "histogram_quantile(0.95, sum by (le, exported_job, http_route) (rate(http_server_request_duration_seconds_bucket[5m])))"),
        ("RabbitMQ ready", "sum by (queue) (rabbitmq_queue_messages_ready)"),
        ("RabbitMQ unacked", "sum by (queue) (rabbitmq_queue_messages_unacked)"),
        ("Consumer count", "sum by (queue) (rabbitmq_queue_consumers)"),
    ]
    add_table(doc, ["Mục tiêu", "PromQL"], [[a, b] for a, b in queries], [2.1, 4.65])
    add_labeled_paragraph(doc, "Cardinality: ", "Không đưa customer id, order id, email, URL thô hoặc correlation id thành metric label. Dùng trace/log cho dữ liệu từng request; metric chỉ nên có chiều hữu hạn như service, route chuẩn hóa, method, status class.")
    doc.add_heading("9.4 Thêm scrape target", level=2)
    add_code(doc, """
scrape_configs:
  - job_name: commerce-otel
    static_configs:
      - targets: [otel-collector:8889]
  - job_name: rabbitmq
    static_configs:
      - targets: [rabbitmq:15692]
  - job_name: custom-exporter
    scrape_interval: 15s
    static_configs:
      - targets: [custom-exporter:9100]
""")
    doc.add_paragraph(
        "Target phải truy cập được từ container Prometheus. `localhost` trong container là chính Prometheus, không phải Windows và cũng không phải service khác."
    )
    doc.add_heading("9.5 Recording rule và alert rule", level=2)
    add_code(doc, """
groups:
  - name: commerce-http
    interval: 30s
    rules:
      - record: commerce:http_requests:rate5m
        expr: sum by (exported_job) (rate(http_server_request_duration_seconds_count[5m]))
      - alert: CommerceTargetDown
        expr: up{job=~"commerce-otel|rabbitmq"} == 0
        for: 2m
        labels:
          severity: warning
        annotations:
          summary: "Commerce scrape target is down"
""")
    doc.add_paragraph(
        "Để dùng rule, mount file vào container và thêm `rule_files` trong `deploy/prometheus.yaml`. Alerting rule chỉ tạo alert state; muốn gửi thông báo theo Prometheus-native cần Alertmanager, hiện chưa có trong Compose. "
        "Có thể dùng Grafana Alerting thay thế cho local/demo."
    )
    doc.add_heading("9.6 Kiểm tra sau khi custom", level=2)
    add_code(doc, """
docker compose --env-file .env config
docker compose --env-file .env exec -T prometheus `
  promtool check config /etc/prometheus/prometheus.yml
docker compose --env-file .env restart prometheus
""")
    doc.add_heading("9.7 Khi query trả về rỗng", level=2)
    for item in [
        "Chạy `up` để xác nhận target còn được scrape; mở `/targets` đọc Last Scrape Error.",
        "Query tên metric không có filter, chuyển sang Table và xem label thực tế trước khi viết selector.",
        "Tạo traffic trong time range đang xem; counter/rate sẽ không có dữ liệu hữu ích nếu chưa có request.",
        "Đối với `rate`, chọn range vector đủ lớn so với scrape interval 10 giây, thường từ 1–5 phút.",
        "Kiểm tra metric có bị đổi tên/đơn vị bởi OTLP-to-Prometheus exporter sau khi nâng version.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("9.8 Quy tắc alert tốt", level=2)
    add_bullet(doc, "Alert theo tác động người dùng hoặc nguy cơ mất dữ liệu, không chỉ theo một spike kỹ thuật ngắn.")
    add_bullet(doc, "Dùng `for` để giảm nhiễu; labels phục vụ routing, annotations giải thích tác động và runbook.")
    add_bullet(doc, "Test trạng thái Pending/Firing/Resolved và trường hợp Prometheus thiếu dữ liệu.")
    add_bullet(doc, "Recording rule nên có tên ổn định, đơn vị rõ và giúp query dashboard/alert dễ đọc hơn.")

    chapter(doc, "10. Grafana: dashboard, Explore, biến, cảnh báo và provisioning")
    doc.add_heading("10.1 Grafana dùng để làm gì", level=2)
    doc.add_paragraph(
        "Grafana là lớp hiển thị và alerting; nó không lưu metric ứng dụng thay Prometheus. Stack provision sẵn data source Prometheus (`http://prometheus:9090`) và dashboard `Commerce overview`, "
        "gồm request rate và p95 latency. Tài khoản admin lấy từ `GRAFANA_ADMIN_USER/PASSWORD` trong `.env`."
    )
    doc.add_heading("10.2 Xem dashboard và điều tra bằng Explore", level=2)
    for item in [
        "Mở http://localhost:3000, đăng nhập, vào Dashboards → Commerce → Commerce overview.",
        "Chọn time range đủ bao phủ lúc test; bật Auto refresh 5s hoặc 10s khi chạy smoke test.",
        "Panel request rate cho biết lưu lượng; p95 cho biết 95% request nhanh hơn ngưỡng hiển thị.",
        "Chọn Explore → Prometheus, chạy `up`, sau đó paste PromQL từ Chương 9.",
        "Dùng Inspect → Query/Data để phân biệt “không có dữ liệu” với lỗi query hoặc lỗi datasource.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("10.3 Tạo panel mới", level=2)
    for item in [
        "Tạo dashboard mới hoặc copy dashboard provisioned sang dashboard editable.",
        "Add visualization, chọn Prometheus, nhập query p95 theo route.",
        "Legend dùng `{{exported_job}} · {{http_route}}`; Unit chọn seconds; Min=0.",
        "Thêm threshold theo SLO nội bộ, ví dụ warning 0,5s và critical 1s; đây là ngưỡng minh họa, phải hiệu chỉnh từ yêu cầu thật.",
        "Đổi time range và tạo traffic để xác nhận series/legend. Lưu dashboard với tên/thu mục rõ ràng.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("10.4 Tạo biến service", level=2)
    add_code(doc, """
Variable name: service
Type: Query
Data source: Prometheus
Query: label_values(http_server_request_duration_seconds_count, exported_job)

# Sau đó lọc panel:
sum by (http_route) (
  rate(http_server_request_duration_seconds_count{exported_job=~"$service"}[5m])
)
""")
    doc.add_paragraph(
        "Bật Multi-value/Include All nếu cần so sánh nhiều service; dùng toán tử regex `=~` thay vì `=` cho multi-value."
    )
    doc.add_heading("10.5 Provisioning và lưu thay đổi", level=2)
    doc.add_paragraph(
        "Repo mount `deploy/grafana/provisioning` và `deploy/grafana/dashboards` ở chế độ read-only. Provider scan mỗi 30 giây. `allowUiUpdates` không được bật, vì vậy cách bền vững là chỉnh dashboard JSON trong repo hoặc export dashboard mới thành JSON rồi review/commit. "
        "Không chỉ sửa trong UI và kỳ vọng thay đổi tồn tại sau recreate container."
    )
    add_code(doc, """
# Sau khi chỉnh provisioning hoặc JSON dashboard
docker compose --env-file .env restart grafana
docker compose --env-file .env logs --tail 100 grafana
""")
    doc.add_heading("10.6 Tạo alert trong Grafana", level=2)
    for item in [
        "Alerting → Alert rules → New alert rule; chọn Prometheus query, ví dụ target down hoặc p95 vượt ngưỡng.",
        "Thêm Reduce/Threshold expression và `for` duration để tránh alert do spike ngắn.",
        "Gắn labels như team, service, severity; annotations mô tả tác động và runbook URL.",
        "Tạo Contact point và Notification policy. Stack local hiện chưa provision kênh gửi thông báo.",
        "Test bằng cách dừng một service có chủ đích, xác nhận Pending → Firing → Resolved, rồi khởi động lại.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("10.7 Dashboard nên bổ sung", level=2)
    add_bullet(doc, "HTTP: request rate, error ratio, p50/p95/p99 theo service/route.")
    add_bullet(doc, "RabbitMQ: ready, unacked, consumer count, publish/deliver/ack rate và dead queue.")
    add_bullet(doc, "Runtime: process CPU/memory, GC pause/heap, thread pool queue nếu exporter cung cấp.")
    add_bullet(doc, "Saga: cần thêm business metrics trong code như checkout started/completed/cancelled và duration; không suy diễn hoàn toàn từ HTTP.")
    doc.add_heading("10.8 Export và review dashboard", level=2)
    for item in [
        "Sau khi thử nghiệm trong UI, export dashboard JSON và bỏ datasource id phụ thuộc instance nếu cần portability.",
        "Đặt `uid` ổn định, title/folder rõ ràng; review PromQL, unit, legend, threshold và time range mặc định.",
        "Lưu JSON dưới `deploy/grafana/dashboards`, restart Grafana và xác nhận provisioner nạp bản mới.",
        "Không commit credential, contact-point secret hoặc URL chứa token trong dashboard JSON.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("10.9 Bố cục dashboard vận hành", level=2)
    add_bullet(doc, "Hàng 1 — Tình trạng: target up, request rate, error ratio, p95/p99.")
    add_bullet(doc, "Hàng 2 — Dependency: RabbitMQ backlog/unacked/consumer, database/Redis health nếu có exporter.")
    add_bullet(doc, "Hàng 3 — Runtime: CPU, memory, GC, thread pool; dùng cùng biến service/environment.")
    add_bullet(doc, "Panel phải có mô tả ngắn và link sang runbook/Jaeger để operator biết bước tiếp theo.")

    chapter(doc, "11. Jaeger: tìm trace, đọc critical path và tùy biến")
    doc.add_heading("11.1 Jaeger dùng để làm gì", level=2)
    doc.add_paragraph(
        "Jaeger lưu và hiển thị distributed trace. Một trace gồm các span HTTP, gRPC và RabbitMQ producer/consumer liên quan. Nó giúp trả lời request chậm ở đâu, lỗi nằm ở hop nào, retry có xảy ra và context có truyền qua event hay không. "
        "Image local hiện là Jaeger 2.21.0, all-in-one/in-memory; restart có thể làm mất trace."
    )
    doc.add_heading("11.2 Tìm và đọc trace", level=2)
    for item in [
        "Tạo traffic trước; mở http://localhost:16686.",
        "Chọn service `Commerce.Gateway`, `Ordering.Api` hoặc service liên quan; chọn Lookback bao phủ thời điểm gửi request.",
        "Lọc Operation nếu danh sách dài; có thể dùng tags/status khi telemetry đã gắn thuộc tính.",
        "Mở trace, đọc từ root span; thanh dài nhất trên đường phụ thuộc thường là critical path.",
        "Mở span để xem duration, status, resource/service, HTTP/gRPC attributes và event/exception. Không kỳ vọng thấy body/token vì không nên thu thập.",
        "Với checkout, tìm producer span `Commerce.Messaging`, rồi consumer span ở Inventory/Payment/Shipping. Async gap là bình thường.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("11.3 Khi không thấy trace", level=2)
    add_table(
        doc,
        ["Triệu chứng", "Kiểm tra", "Cách xử lý"],
        [
            ["Không có service", "Log app và Collector", "Kiểm tra OTLP endpoint/protocol, network, sampling"],
            ["Có HTTP nhưng mất RabbitMQ", "Message headers + consumer log", "Kiểm tra traceparent/tracestate propagation và ActivitySource"],
            ["Trace rời rạc", "Parent/span context", "Đảm bảo consumer khởi tạo Activity từ extracted context"],
            ["Service trùng/khó hiểu", "resource.service.name", "Đặt service name ổn định, thêm environment/instance attributes"],
            ["Trace biến mất sau restart", "Storage backend", "Đây là giới hạn in-memory local; dùng storage bền vững ở production"],
        ],
        [1.75, 2.25, 2.75],
    )
    doc.add_heading("11.4 Tùy biến instrumentation", level=2)
    doc.add_paragraph(
        "Tạo custom span cho nghiệp vụ dài bằng ActivitySource đã đăng ký. Tên span nên ổn định và mô tả thao tác, ví dụ `checkout.validate-cart`, không chứa order id trong tên. "
        "Order id có thể là span attribute nếu chính sách dữ liệu cho phép, nhưng không dùng làm metric label. Gắn status Error và record exception có chọn lọc; không ghi credential/body nhạy cảm."
    )
    add_code(doc, """
using var activity = CommerceTelemetry.ActivitySource.StartActivity(
    "checkout.validate-cart",
    ActivityKind.Internal);

activity?.SetTag("commerce.operation", "checkout");
// Không gắn token, password hoặc dữ liệu cá nhân vào span.
""")
    doc.add_heading("11.5 Tùy biến UI/storage", level=2)
    add_bullet(doc, "Local: có thể cấu hình UI bằng Jaeger UI config, nhưng giữ Compose đơn giản nếu không có nhu cầu cụ thể.")
    add_bullet(doc, "Production: dùng storage bền vững được hỗ trợ, thiết lập retention/index, capacity và bảo vệ UI bằng SSO/reverse proxy.")
    add_bullet(doc, "Không mở cổng ingest/UI công khai. Bật TLS và xác thực giữa Collector với backend.")
    doc.add_heading("11.6 Đọc trace bất đồng bộ qua RabbitMQ", level=2)
    for item in [
        "Bắt đầu từ HTTP root span và ghi lại trace id, correlation id cùng thời điểm checkout.",
        "Tìm producer span của `Commerce.Messaging`; kiểm tra routing key và publish duration/status.",
        "Đi theo consumer span ở service kế tiếp. Khoảng trống giữa producer và consumer là thời gian nằm trong broker/retry, không nhất thiết là mất trace.",
        "Nếu consumer span có parent/link sai, kiểm tra bước extract `traceparent`/`tracestate` trước khi tạo Activity.",
        "Đối chiếu span lỗi với Ready/Unacked/dead queue và log cùng correlation id; Jaeger không thay thế trạng thái broker.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("11.7 Checklist khi đưa Jaeger lên production", level=2)
    add_bullet(doc, "Chọn storage và retention theo tốc độ ingest; kiểm thử truy vấn khi index lớn.")
    add_bullet(doc, "Đặt sampling policy có thể giải thích; ưu tiên giữ trace lỗi/chậm nhưng không vượt ngân sách.")
    add_bullet(doc, "Bảo vệ Query UI/API bằng SSO/RBAC và network policy; mã hóa đường Collector → Jaeger.")
    add_bullet(doc, "Cảnh báo Collector export failure và storage saturation; backup cấu hình, không coi trace là system of record.")

    chapter(doc, "12. Runbook xử lý sự cố và hardening production")
    doc.add_heading("12.1 Runbook: API trả 401", level=2)
    for item in [
        "Xác nhận header đúng dạng `Authorization: Bearer <token>` và token chưa hết hạn.",
        "So issuer/audience ở `.env` với cấu hình tất cả container; thay `.env` cần recreate container để nhận biến mới.",
        "Kiểm tra clock host/container; clock skew hiện 30 giây.",
        "Đăng nhập lại để loại trừ token cũ. Không in token vào log khi debug.",
        "Đọc log Gateway và service đích; cả hai đều validate token.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("12.2 Runbook: API trả 403", level=2)
    for item in [
        "Decode claim offline để xác nhận `role` đúng `Admin`/`Customer`; role claim type hiện là `role`.",
        "Phân biệt role failure với ownership failure. Token hợp lệ nhưng resource không thuộc user vẫn phải bị từ chối.",
        "Kiểm tra user seed và role assignment trong Identity DB; không sửa token thủ công.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("12.3 Runbook: checkout đứng ở trạng thái trung gian", level=2)
    for item in [
        "Lấy order id/correlation id; tìm trace ở Jaeger.",
        "RabbitMQ UI: kiểm tra queue tương ứng Ready, Unacked, Consumers và `.dead`.",
        "Log Outbox publisher/consumer service; xác nhận connection và confirm.",
        "Kiểm tra Outbox/Inbox chỉ trong database service liên quan; không sửa trực tiếp trạng thái saga.",
        "Nếu dead message: sửa nguyên nhân, xác minh idempotency rồi replay có kiểm soát.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("12.4 Runbook: dashboard trống", level=2)
    for item in [
        "Tạo request thực tế và chọn đúng time range.",
        "Prometheus `/targets`: commerce-otel/rabbitmq phải UP.",
        "Prometheus query `up`, rồi query metric gốc; xem labels thực tế trước khi filter.",
        "Grafana data source Save & Test; URL phải là `http://prometheus:9090` trong Docker, không phải localhost.",
        "Collector/app logs: tìm lỗi OTLP export hoặc config pipeline.",
    ]:
        add_numbered(doc, item)
    doc.add_heading("12.5 Runbook: Navicat/Redis client không kết nối", level=2)
    add_table(
        doc,
        ["Tool", "Thiết lập", "Lỗi thường gặp"],
        [
            ["PostgreSQL client", "Host 127.0.0.1; port POSTGRES_PORT; user từ .env; database cụ thể như commerce_identity", "Để Initial Database trống có thể client thử database cùng tên user và báo database không tồn tại"],
            ["Redis client", "Host 127.0.0.1; port REDIS_PORT hoặc 6379; Authentication=None", "Container chưa recreate sau khi thêm ports; port host bị chiếm; mapping chỉ bind loopback"],
        ],
        [1.4, 3.0, 2.35],
    )
    add_code(doc, """
docker compose --env-file .env ps postgres redis
docker compose --env-file .env port postgres 5432
docker compose --env-file .env port redis 6379
Test-NetConnection 127.0.0.1 -Port 5432
Test-NetConnection 127.0.0.1 -Port 6379
""")
    doc.add_heading("12.6 Production hardening theo lớp", level=2)
    hardening_rows = [
        ["Identity/JWT", "OIDC + asymmetric signing/JWKS; key rotation; secret manager; MFA/admin policy"],
        ["Authorization", "Policy + ownership checks; scope/permission rõ; audit security decisions"],
        ["Service-to-service", "mTLS/workload identity; network policy; authenticate gRPC"],
        ["RabbitMQ", "TLS; per-service user/vhost/permission; HA/quorum; alert DLQ/backlog"],
        ["Database/Redis", "Không publish public; TLS/auth; least privilege; backup/restore test"],
        ["Observability", "SSO/RBAC; TLS; retention; redaction; sampling; cardinality budget"],
        ["Resilience", "Timeout, retry có jitter, circuit breaker ở call sync; capacity/load test"],
        ["Delivery", "Không commit `.env`; image pinning/SBOM/scan; migration và rollback có kiểm soát"],
    ]
    add_table(doc, ["Lớp", "Việc cần làm"], hardening_rows, [1.55, 5.2])
    doc.add_heading("12.7 Definition of done cho thay đổi kiến trúc", level=2)
    for item in [
        "Build/test qua; không phá Clean Architecture dependency direction.",
        "Public REST contract, role/ownership và Postman collection đồng bộ.",
        "gRPC `.proto` backward-compatible hoặc được version rõ.",
        "Integration event, routing key, queue, retry/DLQ, Outbox/Inbox được cập nhật và test.",
        "Metric/trace không chứa secret/PII và không tăng cardinality vô hạn.",
        "Dashboard/alert/runbook/documentation cập nhật cùng code.",
    ]:
        add_bullet(doc, item)

    chapter(doc, "Phụ lục A. Lệnh vận hành nhanh")
    doc.add_heading("A.1 Docker Compose", level=2)
    add_code(doc, """
docker compose --env-file .env up -d --build
docker compose --env-file .env ps
docker compose --env-file .env logs -f --tail 200 ordering-api inventory-api payment-api
docker compose --env-file .env restart otel-collector prometheus grafana
docker compose --env-file .env config
""")
    doc.add_heading("A.2 Health và observability", level=2)
    add_code(doc, """
Invoke-WebRequest -UseBasicParsing http://localhost:8080/health/ready
Invoke-RestMethod http://localhost:9090/api/v1/targets
Start-Process http://localhost:15672  # RabbitMQ
Start-Process http://localhost:16686  # Jaeger
Start-Process http://localhost:9090   # Prometheus
Start-Process http://localhost:3000   # Grafana
""")
    doc.add_heading("A.3 Database local", level=2)
    add_code(doc, """
docker compose --env-file .env exec -T postgres `
  psql -U <POSTGRES_USER> -d commerce_ordering -c '\\dt'

docker compose --env-file .env exec -T redis redis-cli ping
""")

    chapter(doc, "Phụ lục B. Danh mục routing key đầy đủ")
    events = [
        ["ordering.order-created.v1", "Ordering", "Notification"],
        ["ordering.inventory-reservation-requested.v1", "Ordering", "Inventory"],
        ["ordering.inventory-release-requested.v1", "Ordering", "Inventory"],
        ["ordering.payment-requested.v1", "Ordering", "Payment"],
        ["ordering.order-paid.v1", "Ordering", "Inventory, Shipping"],
        ["ordering.order-cancelled.v1", "Ordering", "Notification"],
        ["inventory.reserved.v1", "Inventory", "Ordering"],
        ["inventory.reservation-failed.v1", "Inventory", "Ordering"],
        ["inventory.released.v1", "Inventory", "Ordering"],
        ["inventory.reservation-expired.v1", "Inventory", "Ordering"],
        ["payment.succeeded.v1", "Payment", "Ordering, Notification"],
        ["payment.failed.v1", "Payment", "Ordering, Notification"],
        ["shipping.shipment-created.v1", "Shipping", "Ordering, Notification"],
        ["shipping.shipment-delivered.v1", "Shipping", "Ordering, Notification"],
    ]
    add_table(doc, ["Routing key", "Publisher", "Consumer(s)"], events, [3.55, 1.25, 1.95])
    doc.add_heading("B.1 Các queue consumer", level=2)
    queue_names = [
        "inventory.reserve-order.v1 · inventory.release-order.v1 · inventory.confirm-paid-order.v1",
        "payment.process-order.v1",
        "shipping.create-paid-order.v1",
        "ordering.inventory-reserved.v1 · ordering.inventory-failed.v1 · ordering.inventory-expired.v1 · ordering.inventory-released.v1",
        "ordering.payment-succeeded.v1 · ordering.payment-failed.v1 · ordering.shipment-created.v1 · ordering.shipment-delivered.v1",
        "notification.order-created.v1 · notification.order-cancelled.v1 · notification.payment-succeeded.v1 · notification.payment-failed.v1",
        "notification.shipment-created.v1 · notification.shipment-delivered.v1",
    ]
    for name in queue_names:
        add_bullet(doc, name)
    doc.add_paragraph("Mỗi main queue có thêm `<queue>.retry` và `<queue>.dead` theo topology ở Chương 7.")

    chapter(doc, "Phụ lục C. Tệp cấu hình và mã nguồn cần đọc")
    sources = [
        ["Docker topology/secret mapping", "docker-compose.yml; .env.example"],
        ["Gateway/YARP/rate limit", "src/Gateway/Commerce.Gateway/appsettings.json; Program.cs"],
        ["JWT validation", "src/BuildingBlocks/Commerce.BuildingBlocks.Infrastructure/Security"],
        ["RabbitMQ/Outbox/Inbox", "src/BuildingBlocks/Commerce.BuildingBlocks.Infrastructure/Messaging"],
        ["Integration contracts", "src/BuildingBlocks/Commerce.BuildingBlocks.Contracts/IntegrationEvents"],
        ["OTel Collector", "deploy/otel-collector.yaml"],
        ["Prometheus", "deploy/prometheus.yaml"],
        ["Grafana provisioning", "deploy/grafana/provisioning; deploy/grafana/dashboards"],
        ["Kiến trúc hiện có", "docs/architecture.md; docs/local-infrastructure-guide.md"],
    ]
    add_table(doc, ["Chủ đề", "Vị trí"], sources, [2.1, 4.65])

    chapter(doc, "Phụ lục D. Tài liệu tham khảo chính thức")
    references = [
        ("ASP.NET Core JWT bearer authentication", "https://learn.microsoft.com/aspnet/core/security/authentication/configure-jwt-bearer-authentication"),
        ("RabbitMQ Consumer Acknowledgements and Publisher Confirms", "https://www.rabbitmq.com/docs/confirms"),
        ("RabbitMQ Exchanges", "https://www.rabbitmq.com/docs/exchanges"),
        ("RabbitMQ Management Plugin", "https://www.rabbitmq.com/docs/management"),
        ("OpenTelemetry Collector", "https://opentelemetry.io/docs/collector/"),
        ("OpenTelemetry Collector Components", "https://opentelemetry.io/docs/collector/components/"),
        ("Prometheus Querying Basics", "https://prometheus.io/docs/prometheus/latest/querying/basics/"),
        ("Prometheus Recording Rules", "https://prometheus.io/docs/prometheus/latest/configuration/recording_rules/"),
        ("Prometheus Alerting Rules", "https://prometheus.io/docs/prometheus/latest/configuration/alerting_rules/"),
        ("Grafana Dashboards", "https://grafana.com/docs/grafana/latest/fundamentals/dashboards-overview/"),
        ("Grafana Dashboard Variables", "https://grafana.com/docs/grafana/latest/visualizations/dashboards/variables/"),
        ("Grafana Provisioning", "https://grafana.com/tutorials/provision-dashboards-and-data-sources/"),
        ("Jaeger 2.21 Documentation", "https://www.jaegertracing.io/docs/2.21/"),
        ("Jaeger Storage", "https://www.jaegertracing.io/docs/2.21/storage/"),
        ("Jaeger Troubleshooting", "https://www.jaegertracing.io/docs/2.21/operations/troubleshooting/"),
    ]
    for index, (label, url) in enumerate(references, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.25)
        add_text(paragraph, f"{index}. ", bold=True)
        add_hyperlink(paragraph, label, url)

    doc.add_heading("Kết luận", level=2)
    doc.add_paragraph(
        "Commerce tách rõ giao tiếp công khai, truy vấn nội bộ và choreography bất đồng bộ. Độ tin cậy không đến từ một công cụ đơn lẻ mà từ tổ hợp timeout, idempotency, Outbox/Inbox, publisher confirm, retry/DLQ và quan sát xuyên suốt. "
        "Khi custom, hãy thay đổi nhỏ, validate cấu hình, tạo traffic kiểm chứng, rồi đối chiếu cả metric, trace, queue và trạng thái domain."
    )

    doc.core_properties.title = "Sổ tay giao tiếp, bảo mật và quan sát hệ thống Commerce"
    doc.core_properties.subject = "Microservices, JWT, RabbitMQ, OpenTelemetry, Prometheus, Grafana và Jaeger"
    doc.core_properties.author = "Commerce Engineering"
    doc.core_properties.keywords = "Commerce, microservices, JWT, RabbitMQ, OpenTelemetry, Prometheus, Grafana, Jaeger"
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
